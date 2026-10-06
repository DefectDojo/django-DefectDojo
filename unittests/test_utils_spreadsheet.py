import csv
import io

from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from dojo.utils_spreadsheet import TextCellWriter, as_text_cell, store_cells_as_text


class TestSpreadsheetCells(SimpleTestCase):

    def test_as_text_cell(self):
        for value in ("=SUM(A1:A2)", "+1", "-1+1", "@SUM(A1)", "\tx", "\rx"):
            with self.subTest(value=value):
                self.assertEqual("'" + value, as_text_cell(value))
        for value in ("plain", "", "a=b", 42, None, -1):
            with self.subTest(value=value):
                self.assertEqual(value, as_text_cell(value))

    def test_csv_writer(self):
        buffer = io.StringIO()
        writer = TextCellWriter(buffer)
        writer.writerow(["title", '=HYPERLINK("https://example.com")', 3])
        writer.writerows([["ok", "@x", None]])
        rows = list(csv.reader(io.StringIO(buffer.getvalue())))
        self.assertEqual(["title", "'=HYPERLINK(\"https://example.com\")", "3"], rows[0])
        self.assertEqual(["ok", "'@x", ""], rows[1])

    def test_workbook_cells_stored_as_text(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.cell(row=1, column=1, value="=1+1")
        worksheet.cell(row=1, column=2, value="plain")
        worksheet.cell(row=1, column=3, value=7)
        store_cells_as_text(workbook)

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        loaded = load_workbook(buffer).active
        self.assertEqual("s", loaded.cell(row=1, column=1).data_type)
        self.assertEqual("=1+1", loaded.cell(row=1, column=1).value)
        self.assertEqual("plain", loaded.cell(row=1, column=2).value)
        self.assertEqual(7, loaded.cell(row=1, column=3).value)
